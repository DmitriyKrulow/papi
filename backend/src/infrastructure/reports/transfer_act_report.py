# backend/src/infrastructure/reports/transfer_act_report.py
"""Генератор акта передачи имущества в пользование сотруднику"""
from datetime import datetime
from typing import Dict, List, Optional

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter


class TransferActReportGenerator:
    """Генератор акта передачи имущества (Excel)"""

    HEADER_FONT = Font(name='Arial', size=14, bold=True)
    TITLE_FONT = Font(name='Arial', size=10, bold=False)
    SUBTITLE_FONT = Font(name='Arial', size=10)
    HEADER_ROW_FONT = Font(name='Arial', size=10, bold=True)
    SIGNATURE_FONT = Font(name='Arial', size=10)

    HEADER_FILL = PatternFill(start_color='D9D9D9', end_color='D9D9D9', fill_type='solid')

    THIN_BORDER = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin'),
    )

    CENTER_ALIGN = Alignment(horizontal='center', vertical='center', wrap_text=True)
    LEFT_ALIGN = Alignment(horizontal='left', vertical='center', wrap_text=True)

    def __init__(self):
        self.workbook = Workbook()
        self.workbook.remove(self.workbook.active)

    def generate(
        self,
        act_number: str,
        act_date: Optional[datetime],
        employee_name: str = '',
        employee_position: str = '',
        department_name: str = '',
        room_name: str = '',
        assets: List[Dict] = None,
    ) -> bytes:
        """
        Генерация акта передачи имущества.

        Args:
            act_number: номер акта
            act_date: дата акта
            employee_name: ФИО сотрудника-получателя
            employee_position: должность получателя
            department_name: подразделение
            room_name: кабинет
            assets: [{name, inventory_number, quantity, model, serial_number}]

        Returns:
            bytes: содержимое Excel-файла
        """
        assets = assets or []
        sheet = self.workbook.create_sheet('Акт передачи')

        # Ширины колонок
        sheet.column_dimensions['A'].width = 7       # № п.п.
        sheet.column_dimensions['B'].width = 35      # Наименование
        sheet.column_dimensions['C'].width = 32      # Инв. номер
        sheet.column_dimensions['D'].width = 8       # Кол.
        sheet.column_dimensions['E'].width = 22      # Модель
        sheet.column_dimensions['F'].width = 24      # Серийный номер
        sheet.column_dimensions['G'].width = 28      # Подпись

        # Заголовок
        sheet.merge_cells('A2:G2')
        title_cell = sheet['A2']
        title_cell.value = f'АКТ ПЕРЕДАЧИ ИМУЩЕСТВА В ПОЛЬЗОВАНИЕ № {act_number}'
        title_cell.font = self.HEADER_FONT
        title_cell.alignment = Alignment(horizontal='center', vertical='center')

        # Дата
        sheet.merge_cells('A3:G3')
        date_cell = sheet['A3']
        date_cell.value = f'от {act_date.strftime("%d.%m.%Y") if act_date else ""}'
        date_cell.font = self.SUBTITLE_FONT
        date_cell.alignment = Alignment(horizontal='center', vertical='center')

        # Сотрудник-получатель
        sheet.merge_cells('A4:G4')
        emp_cell = sheet['A4']
        parts = []
        if employee_position:
            parts.append(employee_position)
        if employee_name:
            parts.append(employee_name)
        if room_name:
            parts.append(f'каб. {room_name}')
        if department_name:
            parts.append(department_name)
        emp_cell.value = 'Получатель: ' + ', '.join(parts)
        emp_cell.font = self.SUBTITLE_FONT
        emp_cell.alignment = Alignment(horizontal='left', vertical='center')

        # Заголовки таблицы
        headers = ['№ п.п.', 'Наименование', 'Инвентарный номер', 'Кол.', 'Модель', 'Серийный номер', 'Подпись получателя']
        for col, header in enumerate(headers, 1):
            cell = sheet.cell(row=6, column=col, value=header)
            cell.font = self.HEADER_ROW_FONT
            cell.alignment = self.CENTER_ALIGN
            cell.fill = self.HEADER_FILL
            cell.border = self.THIN_BORDER

        # Данные
        for idx, asset in enumerate(assets, start=7):
            row = idx
            num = row - 6

            cell_a = sheet.cell(row=row, column=1, value=num)
            cell_a.font = self.TITLE_FONT
            cell_a.alignment = self.CENTER_ALIGN
            cell_a.border = self.THIN_BORDER

            cell_b = sheet.cell(row=row, column=2, value=asset.get('name', ''))
            cell_b.font = self.TITLE_FONT
            cell_b.alignment = self.LEFT_ALIGN
            cell_b.border = self.THIN_BORDER

            cell_c = sheet.cell(row=row, column=3, value=asset.get('inventory_number', ''))
            cell_c.font = self.TITLE_FONT
            cell_c.alignment = self.CENTER_ALIGN
            cell_c.border = self.THIN_BORDER

            cell_d = sheet.cell(row=row, column=4, value=asset.get('quantity', 1))
            cell_d.font = self.TITLE_FONT
            cell_d.alignment = self.CENTER_ALIGN
            cell_d.border = self.THIN_BORDER

            cell_e = sheet.cell(row=row, column=5, value=asset.get('model', ''))
            cell_e.font = self.TITLE_FONT
            cell_e.alignment = self.LEFT_ALIGN
            cell_e.border = self.THIN_BORDER

            cell_f = sheet.cell(row=row, column=6, value=asset.get('serial_number', ''))
            cell_f.font = self.TITLE_FONT
            cell_f.alignment = self.LEFT_ALIGN
            cell_f.border = self.THIN_BORDER

            cell_g = sheet.cell(row=row, column=7, value='')
            cell_g.border = self.THIN_BORDER

        # Подписи
        sig_row = len(assets) + 9
        sheet.merge_cells(start_row=sig_row, start_column=1, end_row=sig_row, end_column=3)
        sheet.cell(row=sig_row, column=1, value='Передал (должность, ФИО): ______________').font = self.SIGNATURE_FONT
        sheet.merge_cells(start_row=sig_row + 1, start_column=1, end_row=sig_row + 1, end_column=3)
        sheet.cell(row=sig_row + 1, column=1, value='Принял (должность, ФИО): ______________').font = self.SIGNATURE_FONT
        sheet.merge_cells(start_row=sig_row + 2, start_column=1, end_row=sig_row + 2, end_column=3)
        sheet.cell(row=sig_row + 2, column=1, value='Подпись получателя / дата: ____________').font = self.SIGNATURE_FONT

        buf = __import__('io').BytesIO()
        self.workbook.save(buf)
        buf.seek(0)
        return buf.read()