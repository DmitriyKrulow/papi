# backend/src/infrastructure/reports/room_inventory_report.py
from datetime import datetime
from typing import Dict, List, Optional

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter


class RoomInventoryReportGenerator:
    """Генератор описи имущества в помещении"""

    HEADER_FONT = Font(name='Arial', size=14, bold=True)
    TITLE_FONT = Font(name='Arial', size=10, bold=False)
    SUBTITLE_FONT = Font(name='Arial', size=10)
    HEADER_ROW_FONT = Font(name='Arial', size=10)
    SIGNATURE_FONT = Font(name='Arial', size=10)
    SIGNATURE_SUB_FONT = Font(name='Arial', size=8)

    HEADER_FILL = PatternFill(start_color='D9D9D9', end_color='D9D9D9', fill_type='solid')

    THIN_BORDER = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin'),
    )

    CENTER_ALIGN = Alignment(horizontal='center', vertical='center', wrap_text=True)
    LEFT_ALIGN = Alignment(horizontal='left', vertical='center')
    RIGHT_ALIGN = Alignment(horizontal='right', vertical='bottom')
    BOTTOM_ALIGN = Alignment(horizontal='left', vertical='bottom')

    def __init__(self):
        self.workbook = Workbook()
        self.workbook.remove(self.workbook.active)

    def generate(
        self,
        room_name: str,
        department_name: str,
        assets: List[Dict],
        senior_inspector: str = '',
        position: str = '',
    ) -> bytes:
        """
        Генерация описи имущества в помещении.

        Args:
            room_name: Название помещения (кабинет)
            department_name: Подразделение
            assets: Список активов [
                {
                    'name': '...',
                    'inventory_number': '...',
                    'quantity': 1,
                    ...
                }
            ]
            senior_inspector: ФИО старшего инспектора (Фамилия И.О.)
            position: должность старшего инспектора

        Returns:
            bytes: содержимое Excel-файла
        """
        sheet = self.workbook.create_sheet('Опись')

        # --- Колонки: ширина ---
        sheet.column_dimensions['A'].width = 7       # № п.п.
        sheet.column_dimensions['B'].width = 35      # Наименование
        sheet.column_dimensions['C'].width = 32      # Инв. номер
        sheet.column_dimensions['D'].width = 8       # Кол.

        # --- Строка 2: Заголовок ---
        sheet.merge_cells('A2:D2')
        title_cell = sheet['A2']
        title_cell.value = 'ОПИСЬ ИМУЩЕСТВА В ПОМЕЩЕНИИ'
        title_cell.font = self.HEADER_FONT
        title_cell.alignment = Alignment(horizontal='center', vertical='center')

        # --- Строка 3: Кабинет + Подразделение ---
        sheet['B3'] = 'Кабинет №'
        sheet['B3'].font = self.SUBTITLE_FONT
        sheet['B3'].alignment = Alignment(horizontal='right', vertical='bottom')

        room_label = room_name
        if department_name:
            room_label = f'{room_name} ({department_name})'
        sheet['C3'] = room_label
        sheet['C3'].font = self.SUBTITLE_FONT
        sheet['C3'].alignment = Alignment(horizontal='left', vertical='bottom')

        # --- Строка 5: Заголовки таблицы ---
        headers = ['№ п.п.', 'Наименование', 'Инвентарный номер', 'Кол.']
        for col, header in enumerate(headers, 1):
            cell = sheet.cell(row=5, column=col, value=header)
            cell.font = self.HEADER_ROW_FONT
            cell.alignment = self.CENTER_ALIGN
            cell.fill = self.HEADER_FILL
            cell.border = self.THIN_BORDER

        # --- Данные: строки 6+ ---
        for idx, asset in enumerate(assets, start=6):
            row = idx
            num = row - 5  # порядковый номер

            # № п.п.
            cell_a = sheet.cell(row=row, column=1, value=num)
            cell_a.font = self.TITLE_FONT
            cell_a.alignment = self.CENTER_ALIGN
            cell_a.border = self.THIN_BORDER

            # Наименование
            cell_b = sheet.cell(row=row, column=2, value=asset.get('name', ''))
            cell_b.font = self.TITLE_FONT
            cell_b.alignment = self.CENTER_ALIGN
            cell_b.border = self.THIN_BORDER

            # Инв. номер
            cell_c = sheet.cell(row=row, column=3, value=asset.get('inventory_number', ''))
            cell_c.font = self.TITLE_FONT
            cell_c.alignment = self.CENTER_ALIGN
            cell_c.border = self.THIN_BORDER

            # Кол.
            qty = asset.get('quantity', 1)
            if not qty or qty <= 0:
                qty = 1
            cell_d = sheet.cell(row=row, column=4, value=qty)
            cell_d.font = self.TITLE_FONT
            cell_d.alignment = self.CENTER_ALIGN
            cell_d.border = self.THIN_BORDER

        # --- Подписи: строки подписи ---
        last_data_row = 5 + len(assets)
        # Пустые строки между таблицей и подписями
        signature_row = last_data_row + 2

        # Старший инспектор
        sig1 = position or 'Старший инспектор отделения отдела'
        sheet[f'B{signature_row}'] = sig1
        sheet[f'B{signature_row}'].font = self.SIGNATURE_FONT
        sheet[f'B{signature_row}'].alignment = self.BOTTOM_ALIGN

        # Подпись-заглушка
        sig1_sign = '_______________________'
        if senior_inspector:
            sig1_sign = f'{senior_inspector}'
        sheet[f'C{signature_row}'] = sig1_sign
        sheet[f'C{signature_row}'].font = self.SIGNATURE_FONT
        sheet[f'C{signature_row}'].alignment = Alignment(horizontal='left', vertical='bottom')

        # должность / подпись
        sheet[f'B{signature_row + 1}'] = 'должность'
        sheet[f'B{signature_row + 1}'].font = self.SIGNATURE_SUB_FONT
        sheet[f'B{signature_row + 1}'].alignment = Alignment(horizontal='center', vertical='top')

        sheet[f'C{signature_row + 1}'] = 'подпись'
        sheet[f'C{signature_row + 1}'].font = self.SIGNATURE_SUB_FONT
        sheet[f'C{signature_row + 1}'].alignment = Alignment(horizontal='center', vertical='top')

        # --- Фиксированная высота строк (как в шаблоне) ---
        for row_idx in range(2, signature_row + 2):
            sheet.row_dimensions[row_idx].height = 12.8

        return self.save_bytes()

    def save_bytes(self) -> bytes:
        from io import BytesIO
        output = BytesIO()
        self.workbook.save(output)
        return output.getvalue()
